"""Document extraction service for PDF, DOCX, and plain text."""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument

from app.core.config import settings
from app.core.exceptions import ExtractionError
from app.schemas.analysis import DocumentInfo, ExtractionResult, PageText
from app.utils.text_cleaner import clean_text

logger = logging.getLogger(__name__)


class DocumentService:
    """Extracts text content from PDF, DOCX, and plain-text files."""

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def extract(
        self,
        file_bytes: bytes,
        filename: str,
    ) -> ExtractionResult:
        """Route extraction based on file extension."""
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return self._extract_pdf(file_bytes, filename)
        elif ext == ".docx":
            return self._extract_docx(file_bytes, filename)
        elif ext == ".txt":
            return self._extract_text(file_bytes, filename)
        else:
            raise ExtractionError(f"Unsupported extension: {ext}")

    def extract_from_text(self, text: str) -> ExtractionResult:
        """Wrap a plain-text string as an ExtractionResult."""
        cleaned = clean_text(text)
        page = PageText(page=1, text=cleaned)
        doc = DocumentInfo(
            filename="text_input",
            pages=1,
            content_length=len(cleaned),
            extraction_method="plaintext",
        )
        return ExtractionResult(document=doc, pages=[page], full_text=cleaned)

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    def _extract_pdf(self, data: bytes, filename: str) -> ExtractionResult:
        try:
            doc = fitz.open(stream=data, filetype="pdf")
        except Exception as exc:
            raise ExtractionError(f"Cannot open PDF: {exc!s}") from exc

        pages: list[PageText] = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            raw = page.get_text("text")
            cleaned = clean_text(raw)
            pages.append(PageText(page=page_num + 1, text=cleaned))

        doc.close()

        full_text = "\n\n".join(p.text for p in pages if p.text)

        # OCR fallback check
        extraction_method = "pymupdf"
        if len(full_text.strip()) < settings.ocr_min_text_length:
            logger.warning(
                "PDF text below OCR threshold (%d chars). "
                "OCR fallback not yet configured — returning partial text.",
                len(full_text),
            )
            extraction_method = "pymupdf_low_text"

        info = DocumentInfo(
            filename=filename,
            pages=len(pages),
            content_length=len(full_text),
            extraction_method=extraction_method,
        )
        return ExtractionResult(document=info, pages=pages, full_text=full_text)

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------

    def _extract_docx(self, data: bytes, filename: str) -> ExtractionResult:
        try:
            doc = DocxDocument(BytesIO(data))
        except Exception as exc:
            raise ExtractionError(f"Cannot open DOCX: {exc!s}") from exc

        paragraphs: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Also pull text from tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        full_text = clean_text("\n".join(paragraphs))

        # DOCX doesn't have real "pages" — report as 1
        page = PageText(page=1, text=full_text)
        info = DocumentInfo(
            filename=filename,
            pages=1,
            content_length=len(full_text),
            extraction_method="docx",
        )
        return ExtractionResult(document=info, pages=[page], full_text=full_text)

    # ------------------------------------------------------------------
    # Plain text
    # ------------------------------------------------------------------

    def _extract_text(self, data: bytes, filename: str) -> ExtractionResult:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = data.decode("latin-1")
            except Exception as exc:
                raise ExtractionError(f"Cannot decode text file: {exc!s}") from exc

        cleaned = clean_text(text)
        page = PageText(page=1, text=cleaned)
        info = DocumentInfo(
            filename=filename,
            pages=1,
            content_length=len(cleaned),
            extraction_method="plaintext",
        )
        return ExtractionResult(document=info, pages=[page], full_text=cleaned)


# Module-level singleton
document_service = DocumentService()
